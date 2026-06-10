#!/usr/bin/env python3
"""
Compile fill results into source_v2.xlsx
Block 1 (rows 2-361): fill cols F,G,H,I (Niveau, CA, Année, Source)
Block 2 (rows 594-953): also update col C (Secteur) and col A (A garder)

Col indices (1-based openpyxl):
  A=1  B=2  C=3  D=4  E=5  F=6  G=7  H=8  I=9
"""
import os
import re
import openpyxl

FILL_DIR = '/home/user/ma-sourcing-slb/work/fill'
SOURCE = '/home/user/ma-sourcing-slb/work/source_v2.xlsx'

# --- Parse fill files ---

def parse_b1_file(path):
    """Format: NUM_LIGNE | CA | ANNEE | URL | NIVEAU"""
    rows = {}
    if not os.path.exists(path):
        print(f'MISSING: {path}')
        return rows
    with open(path) as f:
        for line in f:
            line = line.strip()
            if not line or not line[0].isdigit():
                continue
            parts = [p.strip() for p in line.split('|')]
            if len(parts) < 5:
                continue
            try:
                num = int(parts[0])
                ca = parts[1]
                annee = parts[2]
                url = parts[3]
                niveau = parts[4]
                rows[num] = {'ca': ca, 'annee': annee, 'url': url, 'niveau': niveau}
            except:
                pass
    return rows

def parse_b2_file(path):
    """Format: NUM_LIGNE | SECTEUR | CA | ANNEE | URL | NIVEAU | A_GARDER"""
    rows = {}
    if not os.path.exists(path):
        print(f'MISSING: {path}')
        return rows
    with open(path) as f:
        for line in f:
            line = line.strip()
            if not line or not line[0].isdigit():
                continue
            parts = [p.strip() for p in line.split('|')]
            if len(parts) < 7:
                # Try 6 fields (missing A_GARDER)
                if len(parts) >= 6:
                    parts.append('n.d.')
                else:
                    continue
            try:
                num = int(parts[0])
                secteur = parts[1]
                ca = parts[2]
                annee = parts[3]
                url = parts[4]
                niveau = parts[5]
                a_garder_raw = parts[6]
                # Normalize A_GARDER
                if 'non' in a_garder_raw.lower() or ca == 'n.d.' or '<200M' in niveau.lower() or 'small' in niveau.lower():
                    a_garder = 'Non'
                else:
                    a_garder = 'Oui'
                rows[num] = {
                    'secteur': secteur,
                    'ca': ca,
                    'annee': annee,
                    'url': url,
                    'niveau': niveau,
                    'a_garder': a_garder,
                }
            except:
                pass
    return rows

# Load all b1 files
b1_files = [
    'b1_datacenter', 'b1_ia', 'b1_quant', 'b1_aero',
    'b1_energie_0', 'b1_energie_1', 'b1_naval', 'b1_crypto'
]
b1_data = {}
for fname in b1_files:
    d = parse_b1_file(os.path.join(FILL_DIR, fname + '.txt'))
    b1_data.update(d)
    print(f'b1 {fname}: {len(d)} rows')

# Load all b2 files
b2_files = [
    'b2_datacenter', 'b2_ia', 'b2_quant', 'b2_aero',
    'b2_naval', 'b2_crypto'
]
b2_data = {}
for fname in b2_files:
    d = parse_b2_file(os.path.join(FILL_DIR, fname + '.txt'))
    b2_data.update(d)
    print(f'b2 {fname}: {len(d)} rows')

print(f'\nTotal b1 rows: {len(b1_data)}')
print(f'Total b2 rows: {len(b2_data)}')

# --- Update Excel ---
wb = openpyxl.load_workbook(SOURCE)
ws = wb['Liste']

changes = 0
skipped_b1 = []
skipped_b2 = []

for row_idx in range(2, ws.max_row + 1):
    a_val = ws.cell(row=row_idx, column=1).value  # A garder

    # Task 1: Block 1 rows 2-361 (A verifier)
    if row_idx <= 361:
        if str(a_val).strip().lower() in ('à vérifier', 'a vérifier', 'a verifier') or a_val is None:
            if row_idx in b1_data:
                d = b1_data[row_idx]
                ws.cell(row=row_idx, column=6).value = d['niveau']  # F
                ws.cell(row=row_idx, column=7).value = d['ca']       # G
                ws.cell(row=row_idx, column=8).value = d['annee']    # H
                ws.cell(row=row_idx, column=9).value = d['url']      # I
                changes += 1
            else:
                skipped_b1.append(row_idx)

    # Task 2: Block 2 rows 594-953
    elif 594 <= row_idx <= 953:
        if row_idx in b2_data:
            d = b2_data[row_idx]
            ws.cell(row=row_idx, column=3).value = d['secteur']    # C
            ws.cell(row=row_idx, column=6).value = d['niveau']     # F
            ws.cell(row=row_idx, column=7).value = d['ca']         # G
            ws.cell(row=row_idx, column=8).value = d['annee']      # H
            ws.cell(row=row_idx, column=9).value = d['url']        # I
            ws.cell(row=row_idx, column=1).value = d['a_garder']   # A
            changes += 1
        else:
            skipped_b2.append(row_idx)

print(f'\nChanges applied: {changes}')
if skipped_b1:
    print(f'B1 rows without fill data ({len(skipped_b1)}): {skipped_b1[:20]}')
if skipped_b2:
    print(f'B2 rows without fill data ({len(skipped_b2)}): {skipped_b2[:20]}')

wb.save(SOURCE)
print(f'\nSaved: {SOURCE}')
